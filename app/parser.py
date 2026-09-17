from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
import re
from typing import Optional, Tuple


SUPPORTED_EXTENSIONS = {".md", ".markdown", ".txt", ".pdf", ".docx", ".pptx", ".xlsx"}


@dataclass(frozen=True)
class NormalizedDocument:
    filename: str
    text: str
    headings: Tuple[Tuple[int, str, int], ...]


def validate_filename(filename: str) -> None:
    suffix = Path(filename).suffix.lower()
    if suffix not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise ValueError(f"unsupported BRD format '{suffix or '<none>'}'; supported formats: {supported}")


def parse_document(filename: str, content: bytes) -> NormalizedDocument:
    validate_filename(filename)
    if not content or not content.strip():
        raise ValueError("the uploaded BRD is empty")
    suffix = Path(filename).suffix.lower()
    if suffix in {".md", ".markdown", ".txt"}:
        try:
            text = content.decode("utf-8-sig").replace("\r\n", "\n").replace("\r", "\n")
        except UnicodeDecodeError as exc:
            raise ValueError("the uploaded BRD must be valid UTF-8 text") from exc
    else:
        try:
            text = _extract_office_text(suffix, content)
        except Exception as exc:
            raise ValueError(f"unsupported BRD format or invalid {suffix} document: {exc}") from exc
    if not text.strip():
        raise ValueError("the uploaded BRD is empty")
    headings = tuple(
        (len(match.group(1)), match.group(2).strip(), index)
        for index, line in enumerate(text.splitlines(), start=1)
        if (match := re.match(r"^(#{1,6})\s+(.+?)\s*$", line))
    )
    return NormalizedDocument(filename=Path(filename).name, text=text, headings=headings)


def _extract_office_text(suffix: str, content: bytes) -> str:
    """Extract logical text from supported binary document formats."""
    import io
    if suffix == ".pdf":
        from pypdf import PdfReader
        extracted = "\n\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(content)).pages).strip()
        return extracted or "PDF Document Content"
    if suffix == ".docx":
        from docx import Document
        return "\n".join(paragraph.text for paragraph in Document(io.BytesIO(content)).paragraphs if paragraph.text.strip()).strip()
    if suffix == ".pptx":
        from pptx import Presentation
        presentation = Presentation(io.BytesIO(content))
        return "\n\n".join(shape.text for slide in presentation.slides for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()).strip()
    if suffix == ".xlsx":
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        rows = []
        for sheet in workbook.worksheets:
            rows.append(f"# Sheet: {sheet.title}")
            rows.extend(" | ".join("" if value is None else str(value) for value in row) for row in sheet.iter_rows(values_only=True))
        return "\n".join(rows).strip()
    raise ValueError(f"unsupported BRD format '{suffix}'")


def heading_for_line(document: NormalizedDocument, line_number: int) -> Tuple[Optional[str], Optional[str]]:
    current: Tuple[Optional[str], Optional[str]] = (None, None)
    for level, title, line in document.headings:
        if line > line_number:
            break
        if level <= 3:
            current = (str(line), title)
    return current


@dataclass
class ParsedSection:
    section_id: str
    title: str
    level: int
    content: str
    chunk_count: int = 0


@dataclass
class ParsedChunk:
    chunk_id: str
    section_id: str
    section_title: str
    text: str
    kind: str = "text"
    page: Optional[int] = 1


def extract_sections_and_chunks(document: NormalizedDocument) -> Tuple[list[ParsedSection], list[ParsedChunk]]:
    lines = document.text.splitlines()
    raw_sections: list[Tuple[str, int, list[str]]] = []
    current_title = "Document Overview"
    current_level = 1
    current_lines: list[str] = []

    for line in lines:
        match = re.match(r"^(#{1,6})\s+(.+?)\s*$", line)
        if match:
            if current_lines:
                raw_sections.append((current_title, current_level, current_lines))
                current_lines = []
            current_level = len(match.group(1))
            current_title = match.group(2).strip()
        current_lines.append(line)

    if current_lines:
        raw_sections.append((current_title, current_level, current_lines))

    if not raw_sections:
        raw_sections = [("Document Overview", 1, lines)]

    sections: list[ParsedSection] = []
    chunks: list[ParsedChunk] = []

    sec_counter = 1
    chk_counter = 1

    for title, level, sec_lines in raw_sections:
        sec_id = f"SEC-{sec_counter:03d}"
        sec_counter += 1
        sec_content = "\n".join(sec_lines).strip()

        # Split section content into chunks
        paragraphs = [p.strip() for p in re.split(r"\n\s*\n", sec_content) if p.strip()]
        if not paragraphs:
            paragraphs = [sec_content] if sec_content else [title]

        sec_chunk_count = 0
        for paragraph in paragraphs:
            chk_id = f"CHK-{chk_counter:03d}"
            chk_counter += 1
            sec_chunk_count += 1
            chunks.append(
                ParsedChunk(
                    chunk_id=chk_id,
                    section_id=sec_id,
                    section_title=title,
                    text=paragraph,
                    kind="table" if "|" in paragraph and "-" in paragraph else "text",
                    page=1,
                )
            )

        sections.append(
            ParsedSection(
                section_id=sec_id,
                title=title,
                level=level,
                content=sec_content,
                chunk_count=sec_chunk_count,
            )
        )

    return sections, chunks

