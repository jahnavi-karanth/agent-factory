from __future__ import annotations

import io
import os
import tempfile
import pytest
from fastapi.testclient import TestClient

from app.main import create_app
from app.repository import SQLiteRepository
from app.document_store import DocumentStore


@pytest.fixture
def temp_repo():
    db_file = tempfile.NamedTemporaryFile(suffix=".sqlite3", delete=False)
    db_file.close()
    repo = SQLiteRepository(path=db_file.name)
    yield repo
    if os.path.exists(db_file.name):
        os.unlink(db_file.name)


@pytest.fixture
def temp_chroma():
    chroma_dir = tempfile.mkdtemp()
    store = DocumentStore(path=chroma_dir)
    yield store


@pytest.fixture
def client(temp_repo, temp_chroma):
    app = create_app(repository=temp_repo, document_store=temp_chroma)
    return TestClient(app)


def get_auth_token(client, email="user1@example.com", password="Password123!"):
    resp = client.post("/auth/register", json={"email": email, "password": password})
    if resp.status_code == 201:
        return resp.json()["access_token"]
    login_resp = client.post("/auth/login", json={"email": email, "password": password})
    return login_resp.json()["access_token"]


# ==========================================
# 1. Authentication Tests
# ==========================================

def test_auth_register_and_login(client):
    reg = client.post("/auth/register", json={"email": "newuser@example.com", "password": "Secret123!"})
    assert reg.status_code == 201
    assert "access_token" in reg.json()
    assert reg.json()["token_type"] == "bearer"

    # Duplicate registration -> 409
    dup = client.post("/auth/register", json={"email": "newuser@example.com", "password": "Secret123!"})
    assert dup.status_code == 409

    # Login via /auth/token
    tok = client.post("/auth/token", json={"email": "newuser@example.com", "password": "Secret123!"})
    assert tok.status_code == 200
    assert "access_token" in tok.json()

    # Login via /auth/login alias
    login_alias = client.post("/auth/login", json={"email": "newuser@example.com", "password": "Secret123!"})
    assert login_alias.status_code == 200
    assert "access_token" in login_alias.json()

    # Wrong password -> 401
    bad = client.post("/auth/login", json={"email": "newuser@example.com", "password": "WrongPassword"})
    assert bad.status_code == 401


def test_protected_routes_without_or_invalid_jwt(client):
    # Missing token -> 401
    r1 = client.get("/projects")
    assert r1.status_code == 401

    # Invalid token -> 401
    r2 = client.get("/projects", headers={"Authorization": "Bearer invalid.jwt.token"})
    assert r2.status_code == 401


# ==========================================
# 2. Project Lifecycle & Ownership Tests
# ==========================================

def test_project_lifecycle_and_status_validation(client):
    token = get_auth_token(client, "proj_owner@example.com")
    headers = {"Authorization": f"Bearer {token}"}

    # Create project (default status: draft)
    p1 = client.post("/projects", json={"name": "Alpha Project"}, headers=headers)
    assert p1.status_code == 201
    proj_data = p1.json()
    assert proj_data["name"] == "Alpha Project"
    assert proj_data["status"] == "draft"
    proj_id = proj_data["project_id"]

    # Invalid status on creation -> 422
    p_bad = client.post("/projects", json={"name": "Bad Proj", "status": "invalid_status"}, headers=headers)
    assert p_bad.status_code == 422

    # Get single project
    get_p = client.get(f"/projects/{proj_id}", headers=headers)
    assert get_p.status_code == 200
    assert get_p.json()["project_id"] == proj_id

    # List projects
    list_p = client.get("/projects", headers=headers)
    assert list_p.status_code == 200
    assert len(list_p.json()) >= 1
    assert any(item["project_id"] == proj_id for item in list_p.json())

    # Update project name & status
    patch_p = client.patch(f"/projects/{proj_id}", json={"name": "Alpha Project V2", "status": "ready"}, headers=headers)
    assert patch_p.status_code == 200
    assert patch_p.json()["name"] == "Alpha Project V2"
    assert patch_p.json()["status"] == "ready"

    # Patch with invalid status -> 422
    patch_bad = client.patch(f"/projects/{proj_id}", json={"status": "unknown_status"}, headers=headers)
    assert patch_bad.status_code == 422


def test_cross_user_project_isolation(client):
    token1 = get_auth_token(client, "user_one@example.com")
    token2 = get_auth_token(client, "user_two@example.com")

    h1 = {"Authorization": f"Bearer {token1}"}
    h2 = {"Authorization": f"Bearer {token2}"}

    proj1_id = client.post("/projects", json={"name": "User 1 Private Proj"}, headers=h1).json()["project_id"]

    # User 2 attempting GET on User 1's project -> 404
    r_get = client.get(f"/projects/{proj1_id}", headers=h2)
    assert r_get.status_code == 404

    # User 2 attempting PATCH on User 1's project -> 404
    r_patch = client.patch(f"/projects/{proj1_id}", json={"name": "Hacked Name"}, headers=h2)
    assert r_patch.status_code == 404


# ==========================================
# 3. Document Upload & File Formats Tests
# ==========================================

def test_document_upload_all_six_formats(client):
    token = get_auth_token(client, "doc_user@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    proj_id = client.post("/projects", json={"name": "MultiFormat Proj"}, headers=headers).json()["project_id"]

    # 1. Markdown (.md)
    md_content = b"# Executive Summary\n\nThis is the markdown document overview.\n\n## System Architecture\n\nDetailed breakdown of components."
    r_md = client.post(f"/projects/{proj_id}/documents", files={"file": ("spec.md", md_content, "text/markdown")}, headers=headers)
    assert r_md.status_code == 201
    doc_md = r_md.json()
    assert doc_md["filename"] == "spec.md"
    assert doc_md["file_type"] == ".md"
    assert doc_md["status"] == "COMPLETED"
    assert doc_md["section_count"] >= 2
    assert doc_md["chunk_count"] >= 2

    # 2. Text (.txt)
    txt_content = b"# Document Title\n\nPlain text content section one.\n\n# Section Two\n\nPlain text content section two."
    r_txt = client.post(f"/projects/{proj_id}/documents", files={"file": ("notes.txt", txt_content, "text/plain")}, headers=headers)
    assert r_txt.status_code == 201
    assert r_txt.json()["file_type"] == ".txt"

    # 3. DOCX (.docx) using python-docx
    from docx import Document
    docx_buf = io.BytesIO()
    doc = Document()
    doc.add_heading("Corporate Expense Management", level=1)
    doc.add_paragraph("Employees shall submit expenses with receipts attached.")
    doc.save(docx_buf)
    r_docx = client.post(f"/projects/{proj_id}/documents", files={"file": ("requirements.docx", docx_buf.getvalue(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")}, headers=headers)
    assert r_docx.status_code == 201
    assert r_docx.json()["file_type"] == ".docx"

    # 4. XLSX (.xlsx) using openpyxl
    from openpyxl import Workbook
    xlsx_buf = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Requirements"
    ws.append(["ID", "Requirement", "Priority"])
    ws.append(["REQ-1", "System must validate user auth", "HIGH"])
    wb.save(xlsx_buf)
    r_xlsx = client.post(f"/projects/{proj_id}/documents", files={"file": ("data.xlsx", xlsx_buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}, headers=headers)
    assert r_xlsx.status_code == 201
    assert r_xlsx.json()["file_type"] == ".xlsx"

    # 5. PPTX (.pptx) using python-pptx
    from pptx import Presentation
    pptx_buf = io.BytesIO()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Project Architecture"
    prs.save(pptx_buf)
    r_pptx = client.post(f"/projects/{proj_id}/documents", files={"file": ("deck.pptx", pptx_buf.getvalue(), "application/vnd.openxmlformats-officedocument.presentationml.presentation")}, headers=headers)
    assert r_pptx.status_code == 201
    assert r_pptx.json()["file_type"] == ".pptx"

    # 6. PDF (.pdf) using pypdf
    from pypdf import PdfWriter
    pdf_buf = io.BytesIO()
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.write(pdf_buf)
    r_pdf = client.post(f"/projects/{proj_id}/documents", files={"file": ("manual.pdf", pdf_buf.getvalue(), "application/pdf")}, headers=headers)
    assert r_pdf.status_code == 201
    assert r_pdf.json()["file_type"] == ".pdf"


def test_unsupported_file_format_and_empty_files(client):
    token = get_auth_token(client, "bad_file_user@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    proj_id = client.post("/projects", json={"name": "Bad File Proj"}, headers=headers).json()["project_id"]

    # Unsupported format -> 400
    r_unsupported = client.post(f"/projects/{proj_id}/documents", files={"file": ("script.py", b"print('hello')", "text/x-python")}, headers=headers)
    assert r_unsupported.status_code == 400

    # Empty file -> 400
    r_empty = client.post(f"/projects/{proj_id}/documents", files={"file": ("empty.txt", b"", "text/plain")}, headers=headers)
    assert r_empty.status_code == 400


def test_path_traversal_filename_sanitization(client):
    token = get_auth_token(client, "security_user@example.com")
    headers = {"Authorization": f"Bearer {token}"}
    proj_id = client.post("/projects", json={"name": "Security Proj"}, headers=headers).json()["project_id"]

    # Path traversal filename
    r_traversal = client.post(f"/projects/{proj_id}/documents", files={"file": ("../../../../etc/passwd.md", b"# Safe Header\n\nContent", "text/markdown")}, headers=headers)
    assert r_traversal.status_code == 201
    doc = r_traversal.json()
    assert doc["filename"] == "passwd.md"
    assert ".." not in doc["storage_path"]


# ==========================================
# 4. Document Status, Sections & Isolation
# ==========================================

def test_document_status_sections_and_cross_project_isolation(client):
    token1 = get_auth_token(client, "user_a@example.com")
    token2 = get_auth_token(client, "user_b@example.com")

    h1 = {"Authorization": f"Bearer {token1}"}
    h2 = {"Authorization": f"Bearer {token2}"}

    proj_a = client.post("/projects", json={"name": "Project A"}, headers=h1).json()["project_id"]
    proj_b = client.post("/projects", json={"name": "Project B"}, headers=h2).json()["project_id"]

    md_a = b"# Heading A1\n\nContent for Project A section 1.\n\n## Heading A2\n\nContent for Project A section 2."
    doc_a = client.post(f"/projects/{proj_a}/documents", files={"file": ("doc_a.md", md_a, "text/markdown")}, headers=h1).json()
    doc_a_id = doc_a["document_id"]

    # Project A owner gets document status -> 200
    status_resp = client.get(f"/projects/{proj_a}/documents/{doc_a_id}", headers=h1)
    assert status_resp.status_code == 200
    assert status_resp.json()["document_id"] == doc_a_id

    # Project A owner gets document sections -> 200
    sections_resp = client.get(f"/projects/{proj_a}/documents/{doc_a_id}/sections", headers=h1)
    assert sections_resp.status_code == 200
    sections = sections_resp.json()
    assert len(sections) >= 2
    assert sections[0]["section_id"].startswith("SEC-")

    # Mismatched Project: User A querying doc_a_id under Project B path -> 404
    mismatch_status = client.get(f"/projects/{proj_b}/documents/{doc_a_id}", headers=h1)
    assert mismatch_status.status_code == 404

    mismatch_sections = client.get(f"/projects/{proj_b}/documents/{doc_a_id}/sections", headers=h1)
    assert mismatch_sections.status_code == 404

    # Unauthorized User B attempting access to Project A doc -> 404
    unauth_status = client.get(f"/projects/{proj_a}/documents/{doc_a_id}", headers=h2)
    assert unauth_status.status_code == 404


# ==========================================
# 5. ChromaDB Embeddings & Search Isolation
# ==========================================

def test_chroma_project_search_isolation(client):
    token1 = get_auth_token(client, "search_a@example.com")
    token2 = get_auth_token(client, "search_b@example.com")

    h1 = {"Authorization": f"Bearer {token1}"}
    h2 = {"Authorization": f"Bearer {token2}"}

    proj_a = client.post("/projects", json={"name": "Project Alpha Search"}, headers=h1).json()["project_id"]
    proj_b = client.post("/projects", json={"name": "Project Beta Search"}, headers=h2).json()["project_id"]

    doc_a_content = b"# Quantum Computing\n\nSuperconducting qubits perform calculations."
    doc_b_content = b"# Renewable Energy\n\nSolar panels generate photovoltaic electricity."

    client.post(f"/projects/{proj_a}/documents", files={"file": ("quantum.md", doc_a_content, "text/markdown")}, headers=h1)
    client.post(f"/projects/{proj_b}/documents", files={"file": ("solar.md", doc_b_content, "text/markdown")}, headers=h2)

    # Search in Project A -> returns quantum content, never solar
    search_a = client.get(f"/projects/{proj_a}/documents/search?q=energy", headers=h1)
    assert search_a.status_code == 200
    results_a = search_a.json()
    for item in results_a:
        assert item["metadata"]["project_id"] == proj_a
        assert "solar" not in item["document"].lower()

    # Search in Project B -> returns solar content, never quantum
    search_b = client.get(f"/projects/{proj_b}/documents/search?q=quantum", headers=h2)
    assert search_b.status_code == 200
    results_b = search_b.json()
    for item in results_b:
        assert item["metadata"]["project_id"] == proj_b
        assert "qubits" not in item["document"].lower()
